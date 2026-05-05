from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl


COST_CATEGORIES = ["equipment", "material", "labor", "land", "catchall"]

STANDARD_COST_COLUMNS = [
    "Total Cost (USD)",
    "Factory Equipment Cost",
    "Site Labor Hours",
    "Site Labor Cost",
    "Site Material Cost",
]

COUNTRY_ALIASES = {
    "korea": "Korea",
    "kor": "Korea",
    "south korea": "Korea",
    "china": "China",
    "chn": "China",
    "uae": "UAE",
    "united arab emirates": "UAE",
}

_EXCEL_PATH = Path(__file__).resolve().parent / "data" / "IAT-v4.1JZcopy.xlsx"

# Tuple indices in each data row (0-based) for LR/SMR Localization sheets
_ACCOUNT_LEVEL_INDICES = [4, 3, 2, 1]  # level 4+, 3, 2, 1 (most-specific first)
_TITLE_IDX = 5
_COUNTRY_START = {"Korea": 11, "China": 16, "UAE": 21}  # first of 5 category cols per country
_CATEGORY_SHARE_INDICES = [33, 34, 35, 36, 37]  # equip%, matl%, labor%, land%, catchall%
_SCENARIO_COST_INDICES = {
    "scenario_1": [39, 40, 41, 42, 43],
    "scenario_2": [44, 45, 46, 47, 48],
    "scenario_3": [49, 50, 51, 52, 53],
}

# Tuple indices for "Results - Adjustment Factors" sheet rows
_AF_COUNTRY_IDX = 1
_AF_KEY_INDICES = {
    "import_tariff": 2,
    "equipment": 3,
    "material": 4,
    "labor": 5,
    "labor_o_and_m": 6,
    "land": 7,
    "catchall": 8,
}


def normalize_account(value: Any) -> str:
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def normalize_country(value: str) -> str:
    key = str(value).strip().lower()
    if key in COUNTRY_ALIASES:
        return COUNTRY_ALIASES[key]
    text = str(value).strip()
    if text in {"Korea", "China", "UAE"}:
        return text
    raise ValueError(f"Unsupported country {value!r}; expected Korea, China, or UAE")


def reactor_family_from_type(reactor_type: str) -> str:
    text = str(reactor_type).strip().lower()
    if "smr" in text:
        return "SMR"
    if "lr" in text or "large" in text or "ap1000" in text:
        return "LR"
    raise ValueError(
        f"Unsupported reactor_type {reactor_type!r}; expected large reactor/LR, SMR, "
        "ACCERT output-LR, or ACCERT output-SMR"
    )


@lru_cache(maxsize=4)
def load_assumptions(path: str | None = None) -> dict[str, Any]:
    excel_path = Path(path) if path else _EXCEL_PATH
    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    return {
        "adjustment_factors": _read_adjustment_factors(wb),
        "families": {
            "LR": _read_family_records(wb, "LR Localization - Inputs"),
            "SMR": _read_family_records(wb, "SMR Localizations - Inputs"),
        },
    }


def _read_adjustment_factors(wb: Any) -> dict[str, dict[str, float]]:
    ws = wb["Results - Adjustment Factors"]
    factors: dict[str, dict[str, float]] = {}
    for row in ws.iter_rows(min_row=4, max_row=6, values_only=True):
        country = row[_AF_COUNTRY_IDX]
        if country is None:
            continue
        factors[str(country)] = {
            key: float(row[idx] or 0.0)
            for key, idx in _AF_KEY_INDICES.items()
        }
    return factors


def _read_family_records(wb: Any, sheet_name: str) -> list[dict[str, Any]]:
    ws = wb[sheet_name]
    records: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=10, values_only=True):
        account_raw = None
        for idx in _ACCOUNT_LEVEL_INDICES:
            if idx < len(row) and row[idx] is not None:
                account_raw = row[idx]
                break
        if account_raw is None:
            continue

        title_raw = row[_TITLE_IDX] if _TITLE_IDX < len(row) else None
        if title_raw is None:
            continue
        title = str(title_raw).strip()
        if not title:
            continue

        account = normalize_account(str(float(account_raw)) if isinstance(account_raw, (int, float)) else str(account_raw))

        localization: dict[str, dict[str, float]] = {}
        for country, start in _COUNTRY_START.items():
            localization[country] = {
                cat: float(row[start + i] or 0.0) if (start + i) < len(row) and row[start + i] is not None else 0.0
                for i, cat in enumerate(COST_CATEGORIES)
            }

        category_shares = {
            cat: float(row[idx] or 0.0) if idx < len(row) and row[idx] is not None else 0.0
            for cat, idx in zip(COST_CATEGORIES, _CATEGORY_SHARE_INDICES)
        }

        reference_costs: dict[str, dict[str, float]] = {}
        for scenario, indices in _SCENARIO_COST_INDICES.items():
            reference_costs[scenario] = {
                cat: float(row[idx] or 0.0) if idx < len(row) and row[idx] is not None else 0.0
                for cat, idx in zip(COST_CATEGORIES, indices)
            }

        records.append(
            {
                "account": account,
                "title": title,
                "localization": localization,
                "category_shares": category_shares,
                "reference_costs_usd_per_kwe": reference_costs,
            }
        )
    return records
