"""Run AP1000 LCOE scenarios used in the EMANES 6 report.

The cases use the local NE-COST AP1000 once-through UOX setup and vary the
capital-cost and O&M assumptions for the three report scenarios:

1. U.S. AP1000 baseline
2. China-calibrated localization and learning case
3. U.S. fast-learning case

Fuel-cycle assumptions are kept consistent with the AP1000 once-through PWR
setup. O&M targets are based on the O&M framework's large-reactor 2023
historical total of 18.81 $/MWh. The China-calibrated case applies only the
China LF_O&M localization factor from Table 3-2 (0.87). The U.S. fast-learning
case retains the U.S. base O&M cost unless a separate O&M learning assumption is
explicitly justified in future work.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import argparse
import shutil
import sys

import pandas as pd
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
SRC = REPO_ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from necostmain import run_necost


OUTPUT_DIR = Path(__file__).with_name("outputs") / "ap1000_report_scenarios"
DEFAULT_IAT_WORKBOOK = Path(
    "/Users/jia.zhou/Library/CloudStorage/Box-Box/EMANES6-FY26/"
    "International Adjustments/IAT-v4.5.xlsx"
)
SAMPLE_SIZE = 1000
OM_TARGET_US_BASELINE = 18.81
OM_FIXED_NOMINAL = 1.9
OM_DIRECT_REFERENCE = 70.0
OM_DIRECT_REFERENCE_MEAN_TOTAL = 10.50743
OM_DIRECT_COMPONENT_PER_UNIT = (OM_DIRECT_REFERENCE_MEAN_TOTAL - OM_FIXED_NOMINAL) / OM_DIRECT_REFERENCE


@dataclass(frozen=True)
class Scenario:
    key: str
    label: str
    occ_per_kw: float
    om_target: float
    description: str

    @property
    def om_direct_nominal(self) -> float:
        return max(1.0, (self.om_target - OM_FIXED_NOMINAL) / OM_DIRECT_COMPONENT_PER_UNIT)


def get_iat_factor(workbook: Path, country: str, factor_name: str) -> float:
    """Read an adjustment factor from IAT's Results - Adjustment Factors sheet."""
    wb = load_workbook(workbook, data_only=True, read_only=True)
    ws = wb["Results - Adjustment Factors"]

    header_row = None
    headers: dict[str, int] = {}
    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        values = [cell.value for cell in row]
        normalized = [str(v).strip().lower() if v is not None else "" for v in values]
        if factor_name.strip().lower() in normalized:
            header_row = row_idx
            headers = {str(v).strip().lower(): idx for idx, v in enumerate(values, start=1) if v is not None}
            break
    if header_row is None:
        raise ValueError(f"Could not find factor header {factor_name!r} in {workbook}")

    factor_col = headers[factor_name.strip().lower()]
    for row_idx in range(header_row + 1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row_idx, col_idx).value
            if isinstance(value, str) and value.strip().lower() == country.strip().lower():
                return float(ws.cell(row_idx, factor_col).value)
    raise ValueError(f"Could not find country {country!r} in {workbook}")


def build_scenarios(iat_workbook: Path) -> list[Scenario]:
    china_om_factor = get_iat_factor(iat_workbook, "China", "Labor O&M")
    return [
        Scenario(
            key="us_baseline",
            label="U.S. AP1000 baseline",
            occ_per_kw=7205.48,
            om_target=OM_TARGET_US_BASELINE,
            description="ACCERT AP1000 baseline OCC with large-reactor O&M framework calibration.",
        ),
        Scenario(
            key="china_calibrated_learning",
            label="China-calibrated learning case",
            occ_per_kw=2989.0,
            om_target=OM_TARGET_US_BASELINE * china_om_factor,
            description=(
                "China IAT localization plus China-calibrated CRF endpoint; "
                f"O&M applies China Labor O&M factor = {china_om_factor:.4f} from IAT-v4.5."
            ),
        ),
        Scenario(
            key="us_fast_learning",
            label="U.S. fast-learning case",
            occ_per_kw=4448.0,
            om_target=OM_TARGET_US_BASELINE,
            description="U.S. cost structure with fast-learning CRF endpoint; U.S. base O&M retained.",
        ),
    ]


SON_TEMPLATE = """necost {{
    construction_interest_rate = 0.05
    operations_interest_rate = 0.05
    sample_size = {sample_size}

    fuel_cycles {{
        cycle(AP1000_REPORT) {{
            reactor(AP1000) {{
                fleet_capacity = 2234
                energy_fraction = 1
            }}
        }}
    }}

    reactors {{
        reactor(AP1000) {{
            capacity_factor = 0.9
            cycle_length = 1.5
            lifetime_years = 60
            power_level {{ reference_thermal = 6.8E9 net_thermal_efficiency = 32.8529 }}
            capital_costs {{ scaling_factor(capital_cost) = 1 }}
            om_costs {{
                scaling_factor(OM_per_year) = 1
                scaling_factor(OM_per_MWh) = 1
            }}
            fuel_reloads {{
                quantity(UOX) {{ heavy_metal_mass = 176.46 fuel_fraction = 1 }}
            }}
        }}
    }}

    capital_costs {{
        item(capital_cost) {{
            cost_type = s_curve
            expenditure_time = 5
            distribution {{ type = triangular low = {occ:.3f} high = {occ:.3f} nominal = {occ:.3f} }}
        }}
    }}

    om_costs {{
        item(OM_per_year) {{ cost_type = fixed distribution {{ type = triangular low = {om_low:.3f} high = {om_high:.3f} nominal = {om_nom:.3f} }} }}
        item(OM_per_MWh) {{ cost_type = variable distribution {{ type = triangular low = 1.2 high = 2.6 nominal = 1.9 }} }}
    }}

    fuel_costs {{
        item(cost_U) {{ cost_value = 110 lead_time = 2 distribution {{ type = triangular low = 65 high = 230 nominal = 110 }} }}
        item(cost_SWU) {{ cost_value = 100 lead_time = 2 distribution {{ type = triangular low = 70 high = 120 nominal = 100 }} }}
        item(cost_fuel_fab) {{ cost_value = 350 lead_time = 0.5 distribution {{ type = triangular low = 200 high = 500 nominal = 350 }} }}
        item(cost_conv) {{ cost_value = 12 lead_time = 2 distribution {{ type = triangular low = 6 high = 18 nominal = 12 }} }}
        item(cost_deconv) {{ cost_value = 6 lead_time = 2 distribution {{ type = triangular low = 4 high = 8 nominal = 6 }} }}
        item(cost_DU_disposal) {{ cost_value = 4 lead_time = 2 distribution {{ type = triangular low = 2 high = 22 nominal = 4 }} }}
    }}

    fuels {{
        fuel(UOX) {{
            avg_discharge_burnup = 50
            num_batches = 3
            avg_specific_power = 1
            fresh_fuel {{
                fabrication {{ loss_fraction = 0.01 costs = [cost_fuel_fab] }}
                EU {{ fuel_fraction = 1 costs = [cost_U] }}
            }}
            EU {{
                conversion {{ loss_fraction = 0.01 costs = [cost_conv cost_deconv] }}
                enrichment {{
                    type = one_stage
                    loss_fraction = 0.01
                    stage_1 {{ feed = 0.711 product = 4.2 tails = 0.25 }}
                    SWU_costs = [cost_SWU]
                    NU_costs = [cost_U]
                    DU_costs = [cost_DU_disposal]
                }}
            }}
        }}
    }}
}}
"""


def scenario_input_text(scenario: Scenario) -> str:
    om_nom = scenario.om_direct_nominal
    return SON_TEMPLATE.format(
        sample_size=SAMPLE_SIZE,
        occ=scenario.occ_per_kw,
        om_low=om_nom * 0.90,
        om_nom=om_nom,
        om_high=om_nom * 1.10,
    )


def summarize(results: pd.DataFrame, scenario: Scenario) -> dict[str, float | str]:
    numeric = results[["Capital", "O&M", "FCC", "LCOE"]]
    row: dict[str, float | str] = {
        "scenario": scenario.label,
        "description": scenario.description,
        "occ_2024_usd_per_kwe": scenario.occ_per_kw,
        "om_target_usd_per_mwh": scenario.om_target,
        "om_direct_spec_usd_per_kwe_year": scenario.om_direct_nominal,
    }
    for col in numeric.columns:
        row[f"{col}_mean"] = numeric[col].mean()
        row[f"{col}_p5"] = numeric[col].quantile(0.05)
        row[f"{col}_p95"] = numeric[col].quantile(0.95)
    return row


def main() -> None:
    parser = argparse.ArgumentParser(description="Run AP1000 NE-COST scenarios for the EMANES 6 Chapter 4 cases.")
    parser.add_argument("--iat-workbook", type=Path, default=DEFAULT_IAT_WORKBOOK, help="IAT workbook used for localization factors.")
    args = parser.parse_args()

    if OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    OUTPUT_DIR.mkdir(parents=True)

    rows = []
    scenarios = build_scenarios(args.iat_workbook)
    for scenario in scenarios:
        scenario_dir = OUTPUT_DIR / scenario.key
        scenario_dir.mkdir()
        input_path = scenario_dir / f"{scenario.key}.son"
        input_path.write_text(scenario_input_text(scenario))
        results = run_necost(input_path, output_dir=scenario_dir, make_plot=False)
        rows.append(summarize(results, scenario))

    summary = pd.DataFrame(rows)
    summary.insert(1, "iat_workbook", str(args.iat_workbook))
    summary_path = OUTPUT_DIR / "ap1000_lcoe_scenario_summary.csv"
    summary.to_csv(summary_path, index=False)

    display_cols = [
        "scenario",
        "occ_2024_usd_per_kwe",
        "Capital_mean",
        "Capital_p5",
        "Capital_p95",
        "O&M_mean",
        "O&M_p5",
        "O&M_p95",
        "FCC_mean",
        "FCC_p5",
        "FCC_p95",
        "LCOE_mean",
        "LCOE_p5",
        "LCOE_p95",
    ]
    print(summary[display_cols].round(2).to_string(index=False))
    print(f"\nScenario outputs written to {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
